"""Builders for small real documents of each supported type, used by extraction tests."""

import json
import zipfile
from pathlib import Path


def pdf(path: Path, pages: list[str], *, title: str = "", author: str = "") -> Path:
    import pymupdf

    doc = pymupdf.open()
    for text in pages:
        page = doc.new_page()
        page.insert_textbox(pymupdf.Rect(36, 36, 559, 806), text, fontsize=9)
    doc.set_metadata({"title": title, "author": author})
    doc.save(path)
    doc.close()
    return path


def docx(path: Path, *, title: str, heading: str, paragraph: str, cells: list[str]) -> Path:
    import docx as python_docx

    document = python_docx.Document()
    document.add_heading(heading, level=1)
    document.add_paragraph(paragraph)
    table = document.add_table(rows=1, cols=len(cells))
    for cell, text in zip(table.rows[0].cells, cells, strict=True):
        cell.text = text
    document.core_properties.title = title
    document.core_properties.keywords = "test, sample"
    document.save(path)
    return path


def pptx(path: Path, *, title: str, slide_title: str, body: str, notes: str) -> Path:
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = slide_title
    slide.placeholders[1].text = body
    table = slide.shapes.add_table(1, 2, Inches(1), Inches(5), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "eigenvalue"
    table.cell(0, 1).text = "variance"
    slide.notes_slide.notes_text_frame.text = notes
    presentation.core_properties.title = title
    presentation.save(path)
    return path


def xlsx(path: Path, *, title: str, sheets: dict[str, list[list[object]]]) -> Path:
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(row)
    workbook.properties.title = title
    workbook.save(path)
    return path


def epub(path: Path, *, title: str, author: str, chapters: list[str]) -> Path:
    from ebooklib import epub as ebook

    book = ebook.EpubBook()
    book.set_identifier("sample-id")
    book.set_title(title)
    book.set_language("en")
    book.add_author(author)
    items = []
    for number, text in enumerate(chapters, start=1):
        chapter = ebook.EpubHtml(title=f"Chapter {number}", file_name=f"ch{number}.xhtml")
        chapter.content = f"<html><body><h1>Chapter {number}</h1><p>{text}</p></body></html>"
        book.add_item(chapter)
        items.append(chapter)
    book.toc = tuple(items)
    book.add_item(ebook.EpubNcx())
    book.add_item(ebook.EpubNav())
    book.spine = items
    ebook.write_epub(str(path), book)
    return path


def notebook(path: Path, cells: list[tuple[str, str]]) -> Path:
    content = {
        "cells": [
            {"cell_type": kind, "metadata": {}, "source": source.splitlines(keepends=True)}
            for kind, source in cells
        ],
        "metadata": {"kernelspec": {"language": "python", "name": "python3"}},
        "nbformat": 4,
        "nbformat_minor": 5,
    }
    path.write_text(json.dumps(content), encoding="utf-8")
    return path


def zip_archive(path: Path, members: list[str]) -> Path:
    with zipfile.ZipFile(path, "w") as archive:
        for name in members:
            archive.writestr(name, "content")
    return path


def mp3(path: Path, *, title: str, artist: str) -> Path:
    from mutagen.id3 import ID3, TIT2, TPE1

    frame = b"\xff\xfb\x90\x64" + bytes(413)  # one silent MPEG-1 Layer III frame
    path.write_bytes(frame * 5)
    tags = ID3()
    tags.add(TIT2(encoding=3, text=title))
    tags.add(TPE1(encoding=3, text=artist))
    tags.save(path)
    return path


def jpeg_with_exif(path: Path, *, description: str, xp_title: str) -> Path:
    from PIL import Image

    exif = Image.Exif()
    exif[0x010E] = description
    exif[0x9C9B] = xp_title.encode("utf-16-le") + bytes(2)
    Image.new("RGB", (8, 8), "white").save(path, exif=exif)
    return path


def png_with_text(path: Path, **text: str) -> Path:
    from PIL import Image
    from PIL.PngImagePlugin import PngInfo

    info = PngInfo()
    for key, value in text.items():
        info.add_itxt(key, value)
    Image.new("RGB", (8, 8), "white").save(path, pnginfo=info)
    return path
