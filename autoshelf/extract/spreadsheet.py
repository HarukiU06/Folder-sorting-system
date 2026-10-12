"""Excel workbooks (.xlsx, .xlsm) via openpyxl, streamed in read-only mode."""

from typing import BinaryIO

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook

from autoshelf.extract.base import Content, TextBudget, non_empty


def extract(file: BinaryIO, max_chars: int) -> Content:
    """Read sheet names and non-empty cell values row by row, plus document properties."""
    workbook = load_workbook(file, read_only=True, data_only=True)
    try:
        budget = TextBudget(max_chars)
        _read_sheets(workbook, budget)
        props = workbook.properties
        return Content(
            text=budget.text,
            title=props.title or None,
            metadata=non_empty(
                author=props.creator,
                subject=props.subject,
                keywords=props.keywords,
                description=props.description,
            ),
            truncated=budget.truncated,
        )
    finally:
        workbook.close()


def _read_sheets(workbook: Workbook, budget: TextBudget) -> None:
    for sheet in workbook.worksheets:
        if not budget.add(sheet.title):
            return
        for row in sheet.iter_rows(values_only=True):
            cells = [str(value).strip() for value in row if value is not None]
            line = " | ".join(cell for cell in cells if cell)
            if line and not budget.add(line):
                return
