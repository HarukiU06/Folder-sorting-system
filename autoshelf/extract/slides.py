"""PowerPoint presentations (.pptx) via python-pptx."""

from collections.abc import Iterable, Iterator
from typing import BinaryIO

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from autoshelf.extract.base import Content, TextBudget, non_empty


def extract(file: BinaryIO, max_chars: int) -> Content:
    """Read slide text (including tables, groups and speaker notes) slide by slide."""
    presentation = Presentation(file)
    budget = TextBudget(max_chars)
    for slide in presentation.slides:
        texts = list(_shape_texts(slide.shapes))
        if slide.has_notes_slide:
            texts.append(slide.notes_slide.notes_text_frame.text)
        if not budget.add("\n".join(texts)):
            break
    props = presentation.core_properties
    return Content(
        text=budget.text,
        title=props.title or None,
        metadata=non_empty(author=props.author, subject=props.subject, keywords=props.keywords),
        truncated=budget.truncated,
    )


def _shape_texts(shapes: Iterable) -> Iterator[str]:
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _shape_texts(shape.shapes)
        elif shape.has_text_frame:
            yield shape.text_frame.text
        elif getattr(shape, "has_table", False):
            for row in shape.table.rows:
                yield " | ".join(cell.text for cell in row.cells if cell.text.strip())
